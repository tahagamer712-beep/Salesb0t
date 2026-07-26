import io
from collections import Counter
from datetime import date, datetime

try:
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ── تحويل الأرقام العربية إلى إنجليزية ─────────────────────────────────────
_AR_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')


def _to_latin(s):
    return str(s).strip().translate(_AR_DIGITS)


def _to_num(val):
    if val is None:
        return None
    s = _to_latin(val)
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except Exception:
        return None


def _parse_date_str(val, default_month, default_year):
    if not val:
        return f"{default_year}-{default_month:02d}-01"
    # ── كائن datetime من openpyxl (Excel يحوّل الخلية لتاريخ تلقائياً) ────────
    # مشكلة: Excel يقرأ "١٢/٤" كـ "ديسمبر 4" (MM/DD) بدل "أبريل 12" (DD/MM)
    # الحل: إذا val.month ≠ الشهر المختار لكن val.day = الشهر المختار → نعكسهما
    if hasattr(val, 'month') and hasattr(val, 'day'):
        try:
            day   = val.day
            month = val.month
            # Excel قرأ بصيغة MM/DD بالخطأ → نعكس
            if month != default_month and day == default_month and 1 <= month <= 31:
                day, month = month, day
            # تحقق من صحة النتيجة
            if not (1 <= day <= 31 and 1 <= month <= 12):
                day, month = 1, default_month
            return date(default_year, month, day).isoformat()
        except Exception:
            return f"{default_year}-{default_month:02d}-01"
    s = _to_latin(str(val)).strip()
    # استبدال كل الفواصل الممكنة (/ و \ و - و |) بـ /
    s = s.replace("\\", "/").replace("-", "/").replace("|", "/")
    # حذف جزء الوقت إن وُجد (مثل "2026-04-01 00:00:00")
    s = s.split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d/%m"):
        try:
            d = datetime.strptime(s, fmt)
            # دائماً نستخدم السنة من اختيار المستخدم
            d = d.replace(year=default_year)
            return d.date().isoformat()
        except Exception:
            pass
    return f"{default_year}-{default_month:02d}-01"


def generate_import_template() -> io.BytesIO:
    """ينشئ قالب Excel فارغ للمندوب لتعبئة بياناته ثم إرساله للبوت."""
    from openpyxl import Workbook

    wb = Workbook()

    # ── الشيت الأولى: البيانات ─────────────────────────────────────────────
    ws = wb.active
    ws.title = "البيانات"
    ws.sheet_view.rightToLeft = True

    headers = ["اسم الصيدلية", "المنطقة", "التاريخ",
               "اسم المنتج", "الكمية", "السعر", "البونص"]
    hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF", size=12)
    ex_fill  = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    examples = [
        ["صيدلية الشفاء",   "الكرادة",  "1/4", "Maxiflex", 10, 18000, 0],
        ["صيدلية النور",    "المنصور",  "2/4", "Ezora",     5,  2500, 2],
        ["صيدلية الأمل",    "الكرادة",  "3/4", "Trazola",   8, 18000, 0],
    ]
    for r_idx, row_data in enumerate(examples, 2):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = ex_fill
            cell.alignment = Alignment(horizontal="center")

    widths = [24, 16, 13, 22, 10, 13, 10]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 22

    # ── الشيت الثانية: تعليمات ──────────────────────────────────────────────
    ws2 = wb.create_sheet("تعليمات")
    ws2.sheet_view.rightToLeft = True
    instructions = [
        ("📋 تعليمات تعبئة ملف الاستيراد", True),
        ("", False),
        ("1. الأعمدة المطلوبة:", True),
        ("   • اسم الصيدلية  — اسم الصيدلية كاملاً (مطلوب)", False),
        ("   • المنطقة       — المنطقة أو الحي (اختياري)", False),
        ("   • التاريخ       — مثال: 1/4  أو  01/04/2025", False),
        ("   • اسم المنتج   — اسم الدواء (مطلوب)", False),
        ("   • الكمية        — رقم فقط مثل: 10  أو  ١٠ (مطلوب)", False),
        ("   • السعر         — سعر الوحدة بالدينار (اختياري)", False),
        ("   • البونص        — كمية البونص، 0 إذا لا يوجد (اختياري)", False),
        ("", False),
        ("2. ملاحظات مهمة:", True),
        ("   • الأرقام: مقبول عربي (١٠) أو إنجليزي (10) أو مختلط", False),
        ("   • كل سطر = طلب واحد (صيدلية واحدة + منتج واحد)", False),
        ("   • إذا صيدلية باعت 3 منتجات → اكتبها في 3 أسطر", False),
        ("   • لا تغير عناوين الأعمدة في السطر الأول الأزرق", False),
        ("   • الأسطر الفارغة تُتجاهل تلقائياً", False),
        ("", False),
        ("3. بعد التعبئة:", True),
        ("   • احفظ الملف وأرسله للبوت مباشرةً", False),
        ("   • البوت سيسألك عن الشهر ثم يعرض معاينة", False),
        ("   • بعد تأكيدك يُحفظ كل شيء في قاعدة البيانات", False),
    ]
    for r_idx, (text, bold) in enumerate(instructions, 1):
        cell = ws2.cell(row=r_idx, column=1, value=text)
        cell.font = Font(bold=bold, size=11)
        cell.alignment = Alignment(horizontal="right")
    ws2.column_dimensions["A"].width = 58

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_catalog_import_template(
    catalog: list,
    rep_name: str,
    month: int,
    year: int,
) -> io.BytesIO:
    """
    ينشئ قالب Excel للاستيراد بنفس شكل قالب التقرير:
    - أعمدة الأدوية من كتالوج المندوب (بدون تاركت)
    - كل سطر = صيدلية واحدة
    - المندوب يكتب الكميات مباشرة في الخلايا
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "البيانات"
    ws.sheet_view.rightToLeft = True

    months_ar_list = ["", "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
                      "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
    month_name = months_ar_list[month] if 1 <= month <= 12 else str(month)

    drug_list  = [d["drug_name"] for d in catalog] if catalog else []
    price_map  = {d["drug_name"]: d.get("price") or 0 for d in catalog} if catalog else {}
    n = len(drug_list)
    total_col = 4 + n

    fill_dark = PatternFill("solid", fgColor="1F3864")
    fill_mid  = PatternFill("solid", fgColor="2E75B6")
    fill_org  = PatternFill("solid", fgColor="FCE4D6")
    fill_ex   = PatternFill("solid", fgColor="EBF3FB")
    fill_row  = PatternFill("solid", fgColor="DEEAF1")

    fw  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    fb  = Font(name="Calibri", bold=True, size=10)
    fn  = Font(name="Calibri", size=10)
    cen = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rgt = Alignment(horizontal="right",  vertical="center")
    thn = Side(style="thin", color="AAAAAA")
    brd = Border(left=thn, right=thn, top=thn, bottom=thn)

    def S(cell, fill=None, font=None, align=None, fmt=None):
        if fill:  cell.fill = fill
        if font:  cell.font = font
        if align: cell.alignment = align
        if fmt:   cell.number_format = fmt
        cell.border = brd

    # ── Row 1: Title ──────────────────────────────────────────────────────────
    if n > 0:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_col)
    S(ws.cell(1, 1, f"📥 قالب الاستيراد — {rep_name} — {month_name} {year}"),
      fill_dark, fw, cen)
    ws.row_dimensions[1].height = 24

    # ── Row 2: ITEMS + drug names ─────────────────────────────────────────────
    S(ws.cell(2, 1, "ITEMS"), fill_mid, fw, cen)
    S(ws.cell(2, 2, ""), fill_mid, fw, cen)
    S(ws.cell(2, 3, ""), fill_mid, fw, cen)
    for i, d in enumerate(drug_list):
        S(ws.cell(2, 4 + i, d), fill_mid, fw, cen)
    if n > 0:
        S(ws.cell(2, total_col, "إجمالي"), fill_mid, fw, cen)
    ws.row_dimensions[2].height = 20

    # ── Row 3: price ──────────────────────────────────────────────────────────
    S(ws.cell(3, 1, "السعر (IQD)"), fill_org, fb, cen)
    S(ws.cell(3, 2, ""), fill_org, fb, cen)
    S(ws.cell(3, 3, ""), fill_org, fb, cen)
    for i, d in enumerate(drug_list):
        p = price_map.get(d, 0)
        S(ws.cell(3, 4 + i, p if p else ""), fill_org, fn, cen, "#,##0")
    if n > 0:
        S(ws.cell(3, total_col, ""), fill_org, fb, cen)
    ws.row_dimensions[3].height = 18

    # ── Row 4: Column headers (MARKER ROW — used by parser to detect catalog format)
    for col, h in enumerate(["الصيدلية", "المنطقة", "التاريخ"], 1):
        S(ws.cell(4, col, h), fill_dark, fw, cen)
    for i, d in enumerate(drug_list):
        S(ws.cell(4, 4 + i, d), fill_dark, fw, cen)
    if n > 0:
        S(ws.cell(4, total_col, "إجمالي"), fill_dark, fw, cen)
    ws.row_dimensions[4].height = 20

    # helper: build إجمالي SUM formula for a given row
    def _sum_formula(r: int) -> str:
        if n == 0:
            return ""
        first = get_column_letter(4)
        last  = get_column_letter(3 + n)
        return f"=SUM({first}{r}:{last}{r})"

    # ── Rows 5–7: Example rows (highlighted) ─────────────────────────────────
    examples = [
        ("صيدلية الشفاء", "الكرادة",  f"1/{month}"),
        ("صيدلية النور",  "المنصور",  f"5/{month}"),
        ("صيدلية الأمل",  "المنصور",  f"10/{month}"),
    ]
    for r_i, (ph, ar, dt) in enumerate(examples):
        row = 5 + r_i
        S(ws.cell(row, 1, ph), fill_ex, fn, rgt)
        S(ws.cell(row, 2, ar), fill_ex, fn, cen)
        S(ws.cell(row, 3, dt), fill_ex, fn, cen)
        for i in range(n):
            S(ws.cell(row, 4 + i, ""), fill_ex, fn, cen)
        if n > 0:
            c = ws.cell(row, total_col, _sum_formula(row))
            S(c, fill_ex, fb, cen, "#,##0")

    # ── Rows 8–200: Empty rows for user data (مع صيغة الإجمالي التلقائية) ────
    for row in range(8, 201):
        rf = fill_row if row % 2 == 0 else None
        S(ws.cell(row, 1, ""), rf, fn, rgt)
        S(ws.cell(row, 2, ""), rf, fn, cen)
        S(ws.cell(row, 3, ""), rf, fn, cen)
        for i in range(n):
            S(ws.cell(row, 4 + i, ""), rf, fn, cen)
        if n > 0:
            c = ws.cell(row, total_col, _sum_formula(row))
            S(c, rf, fb, cen, "#,##0")

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    for i in range(n):
        ws.column_dimensions[get_column_letter(4 + i)].width = 12
    if n > 0:
        ws.column_dimensions[get_column_letter(total_col)].width = 10
    ws.freeze_panes = ws.cell(5, 4)

    # ── Instructions sheet ────────────────────────────────────────────────────
    ws2 = wb.create_sheet("تعليمات")
    ws2.sheet_view.rightToLeft = True
    instructions = [
        ("📋 تعليمات تعبئة قالب الاستيراد", True),
        ("", False),
        (f"✅ هذا القالب مخصص لك — يحوي أدويتك بأسعارها لشهر {month_name} {year}", False),
        ("", False),
        ("طريقة التعبئة:", True),
        ("   • الصيدلية    — اكتب اسم الصيدلية (مطلوب)", False),
        ("   • المنطقة     — المنطقة أو الحي (اختياري)", False),
        ("   • التاريخ     — مثال: 1/4  أو  15/4  (اليوم/الشهر)", False),
        ("   • أعمدة الأدوية — اكتب الكمية في خلية كل دواء", False),
        ("   • اترك الخلية فارغة إذا لم تبع هذا الدواء في هذه الزيارة", False),
        ("", False),
        ("ملاحظات:", True),
        ("   • كل سطر = زيارة صيدلية واحدة", False),
        ("   • الأرقام: مقبول عربي (١٠) أو إنجليزي (10)", False),
        ("   • الأسطر الملوّنة بالأزرق الفاتح (5-7): أمثلة — احذفها أو استبدلها", False),
        ("   • لا تغيّر الصف رقم 4 (الترويسة الداكنة) ولا تضيف/تحذف أعمدة", False),
        ("", False),
        ("بعد التعبئة:", True),
        ("   • احفظ الملف وأرسله للبوت — سيُعالَج تلقائياً", False),
    ]
    for r_idx, (text, bold) in enumerate(instructions, 1):
        cell = ws2.cell(row=r_idx, column=1, value=text)
        cell.font = Font(bold=bold, size=11)
        cell.alignment = Alignment(horizontal="right")
    ws2.column_dimensions["A"].width = 65

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def parse_catalog_import_excel(
    file_bytes: bytes,
    default_month: int,
    default_year: int,
    rep_name: str,
) -> list:
    """
    يقرأ Excel بصيغة قالب الكتالوج (أعمدة أدوية) ويحوّله إلى قائمة طلبات.
    كل سطر = صيدلية واحدة، وكل عمود دواء بكمية > 0 = طلب منفصل.
    يعود بقائمة فارغة إذا لم يتعرّف على الصيغة (للرجوع إلى المحلّل القديم).
    """
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb["البيانات"] if "البيانات" in wb.sheetnames else wb.active

    header_row_idx = None
    drug_col_map   = {}   # drug_name -> col_index (1-based)
    price_map      = {}   # drug_name -> price
    price_row_data = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if not any(c for c in row if c is not None):
            continue
        col1 = _to_latin(str(row[0] or "")).strip().lower()

        # Price row
        if col1 in ("السعر (iqd)", "السعر", "price", "سعر (iqd)", "سعر"):
            price_row_data = row

        # Header row — must have "الصيدلية" / "pharmacy" in col 1
        if col1 in ("الصيدلية", "pharmacy", "صيدلية"):
            header_row_idx = row_idx
            skip_labels = {"إجمالي", "total", "api", "%", ""}
            for col_idx, cell_val in enumerate(row[3:], 4):
                if cell_val and str(cell_val).strip() not in skip_labels:
                    drug_col_map[str(cell_val).strip()] = col_idx
            break

    if not drug_col_map or header_row_idx is None:
        return []   # ليس صيغة الكتالوج

    # Build price_map from the price row
    if price_row_data is not None:
        for drug_name, col_idx in drug_col_map.items():
            val = price_row_data[col_idx - 1] if col_idx - 1 < len(price_row_data) else None
            p = _to_num(val)
            if p:
                price_map[drug_name] = p

    orders = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not any(c for c in row if c is not None):
            continue

        pharmacy = str(row[0] or "").strip()
        if not pharmacy:
            continue

        area     = str(row[1] or "").strip() or None
        date_str = _parse_date_str(
            row[2] if len(row) > 2 else None,
            default_month, default_year,
        )

        # الشهر والسنة دائماً من اختيار المستخدم — كل البيانات في القالب تنتمي للشهر المختار
        # حتى لو التاريخ المكتوب من الشهر السابق (بيانات أواخر الشهر السابق = جزء من دورة الشهر الحالي)
        o_month = default_month
        o_year  = default_year

        for drug_name, col_idx in drug_col_map.items():
            val = row[col_idx - 1] if col_idx - 1 < len(row) else None
            qty = _to_num(val)
            if not qty or qty <= 0:
                continue

            orders.append({
                "rep_name":   rep_name,
                "order_date": date_str,
                "month":      o_month,
                "year":       o_year,
                "area":       area,
                "pharmacy":   pharmacy,
                "item_name":  drug_name,
                "quantity":   int(qty),
                "price":      price_map.get(drug_name, 0),
                "bonus":      0,
                "raw_text":   f"[مستورد] {pharmacy} - {drug_name} x{int(qty)}",
            })

    return orders


def parse_import_excel(file_bytes: bytes, default_month: int, default_year: int,
                       rep_name: str) -> list:
    """
    يقرأ Excel المندوب (بالتنسيق الموحّد) ويحوّله إلى قائمة طلبات.
    كل سطر = (صيدلية + منتج + كمية) → طلب واحد.
    """
    import openpyxl

    HEADER_MAP = {
        "اسم الصيدلية": ["اسم الصيدلية", "الصيدلية", "pharmacy", "صيدلية"],
        "المنطقة":       ["المنطقة", "منطقة", "area", "الحي", "حي"],
        "التاريخ":        ["التاريخ", "date", "تاريخ"],
        "اسم المنتج":    ["اسم المنتج", "المنتج", "دواء", "product", "item", "الدواء", "ايتم"],
        "الكمية":         ["الكمية", "كمية", "qty", "quantity", "عدد"],
        "السعر":          ["السعر", "سعر", "price"],
        "البونص":         ["البونص", "بونص", "bonus"],
    }

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    # نفضّل شيت اسمها "البيانات" وإلا أول شيت
    ws = wb["البيانات"] if "البيانات" in wb.sheetnames else wb.active

    col_map = {}   # canonical_name -> col_index (0-based)
    header_found = False
    orders = []

    for row in ws.iter_rows(values_only=True):
        # تجاهل الأسطر الفارغة كاملاً
        if not any(c for c in row if c is not None):
            continue

        if not header_found:
            row_strs = [_to_latin(c).strip().lower() if c else "" for c in row]
            for canonical, aliases in HEADER_MAP.items():
                for idx, cell_str in enumerate(row_strs):
                    if any(a.lower() in cell_str or cell_str in a.lower()
                           for a in aliases):
                        col_map[canonical] = idx
                        break
            if "اسم الصيدلية" in col_map and "اسم المنتج" in col_map and "الكمية" in col_map:
                header_found = True
            continue

        def get(name):
            idx = col_map.get(name)
            if idx is None or idx >= len(row):
                return None
            v = row[idx]
            # تجاهل كائنات ArrayFormula
            if hasattr(v, "text"):
                return None
            return v

        pharmacy = str(get("اسم الصيدلية") or "").strip()
        product  = str(get("اسم المنتج") or "").strip()
        qty      = _to_num(get("الكمية"))

        if not pharmacy or not product or not qty:
            continue

        area      = str(get("المنطقة") or "").strip() or None
        date_str  = _parse_date_str(get("التاريخ"), default_month, default_year)
        price     = _to_num(get("السعر")) or 0
        bonus     = _to_num(get("البونص")) or 0

        try:
            d = datetime.strptime(date_str, "%Y-%m-%d").date()
            o_month, o_year = d.month, d.year
        except Exception:
            o_month, o_year = default_month, default_year

        orders.append({
            "rep_name":   rep_name,
            "order_date": date_str,
            "month":      o_month,
            "year":       o_year,
            "area":       area,
            "pharmacy":   pharmacy,
            "item_name":  product,
            "quantity":   int(qty),
            "price":      price,
            "bonus":      int(bonus) if bonus else 0,
            "raw_text":   f"[مستورد] {pharmacy} - {product} x{int(qty)}",
        })

    return orders


def _safe_int(value, default=0):
    """بديل pandas.to_numeric(...).fillna(0).astype(int) — يرجع int دائماً، وإلا default."""
    if value is None:
        return default
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return default


def generate_excel_report(orders: list, filters: dict = None) -> io.BytesIO:
    from openpyxl import Workbook
    if not orders:
        return None

    fields = ["rep_name", "order_date", "area", "pharmacy", "doctor_name", "item_name",
              "quantity", "bonus", "price", "payment_type", "notes", "created_at"]
    rows = []
    for o in orders:
        row = {f: o.get(f) for f in fields}
        row["quantity"] = _safe_int(row["quantity"])
        row["bonus"] = _safe_int(row["bonus"])
        rows.append(row)

    display_cols = {
        "rep_name": "اسم المندوب",
        "order_date": "التاريخ",
        "area": "المنطقة",
        "pharmacy": "الصيدلية",
        "doctor_name": "اسم الطبيب",
        "item_name": "المنتج",
        "quantity": "الكمية",
        "bonus": "البونص",
        "price": "السعر",
        "payment_type": "نوع الدفع",
        "notes": "ملاحظات",
        "created_at": "وقت الإدخال",
    }

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "الطلبات"
    ws1.append(list(display_cols.values()))
    for row in rows:
        ws1.append([row[c] for c in display_cols])

    def most_common(field):
        """يطابق df[field].value_counts().idxmax() — أكثر قيمة غير فارغة تكراراً."""
        counter = Counter(r[field] for r in rows if r.get(field) not in (None, ""))
        if not counter:
            return "-"
        return counter.most_common(1)[0][0]

    total_qty = sum(r["quantity"] for r in rows)
    total_bonus = sum(r["bonus"] for r in rows)

    ws2 = wb.create_sheet("الملخص")
    ws2.append(["البيان", "القيمة"])
    for label, value in [
        ("إجمالي الطلبات", len(rows)),
        ("إجمالي الكمية", total_qty),
        ("إجمالي البونص", total_bonus),
        ("أكثر صيدلية", most_common("pharmacy")),
        ("أكثر منتج مبيعاً", most_common("item_name")),
        ("أكثر منطقة", most_common("area")),
        ("أكثر نوع دفع", most_common("payment_type")),
    ]:
        ws2.append([label, value])

    def build_ranking_sheet(sheet_name, field, header_label, sort_by):
        """بديل df.groupby(field).agg(count, sum).sort_values().reset_index() —
        groupby بـpandas الافتراضي يتجاهل صفوف None بالحقل، فنطابق نفس السلوك."""
        agg = {}
        for r in rows:
            key = r.get(field)
            if key in (None, ""):
                continue
            if key not in agg:
                agg[key] = {"count": 0, "qty": 0}
            agg[key]["count"] += 1
            agg[key]["qty"] += r["quantity"]
        if not agg:
            return
        items = list(agg.items())
        if sort_by == "count":
            items.sort(key=lambda kv: kv[1]["count"], reverse=True)
        else:
            items.sort(key=lambda kv: kv[1]["qty"], reverse=True)
        ws = wb.create_sheet(sheet_name)
        ws.append([header_label, "عدد_الطلبات", "إجمالي_الكمية"])
        for key, agg_vals in items:
            ws.append([key, agg_vals["count"], agg_vals["qty"]])

    build_ranking_sheet("ترتيب الصيدليات", "pharmacy", "الصيدلية", sort_by="count")
    build_ranking_sheet("ترتيب المناطق", "area", "المنطقة", sort_by="count")
    build_ranking_sheet("ترتيب المنتجات", "item_name", "المنتج", sort_by="qty")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _trend_icon(curr: int, prev: int) -> str:
    if prev == 0:
        return "🆕"
    pct = (curr - prev) / prev * 100
    if pct >= 20:
        return f"📈 +{pct:.0f}%"
    elif pct <= -20:
        return f"📉 {pct:.0f}%"
    elif pct > 0:
        return f"↗️ +{pct:.0f}%"
    else:
        return f"↘️ {pct:.0f}%"


def build_text_summary(orders: list, report_type: str = "general",
                       prev_orders: list = None) -> str:
    if not orders:
        return "لا توجد بيانات للفترة المحددة."

    quantities = [_safe_int(o.get("quantity", 0)) for o in orders]
    bonuses = [_safe_int(o.get("bonus", 0)) for o in orders]
    total_qty = sum(quantities)
    total_bonus = sum(bonuses)

    trend = ""
    if prev_orders is not None:
        prev_cnt = len(prev_orders) if prev_orders else 0
        trend = f"  {_trend_icon(len(orders), prev_cnt)}"

    lines = [
        f"📊 *إجمالي الطلبات:* {len(orders)}{trend}",
        f"📦 *إجمالي الكمية:* {total_qty}",
        f"🎁 *إجمالي البونص:* {total_bonus}",
    ]

    prices = [o.get("price") for o in orders if o.get("price") not in (None, "")]
    if prices:
        total_price = sum(_safe_int(p) for p in prices)
        if total_price > 0:
            lines.append(f"💵 *إجمالي المبلغ:* {total_price:,.0f}")

    def top_counts(field, n):
        counter = Counter(o.get(field) for o in orders if o.get(field) not in (None, ""))
        return counter.most_common(n)

    if report_type in ("general", "pharmacies"):
        top = top_counts("pharmacy", 3)
        if top:
            lines.append("\n🏆 *أفضل 3 صيدليات:*")
            medals = ["🥇", "🥈", "🥉"]
            for i, (ph, cnt) in enumerate(top):
                lines.append(f"  {medals[i]} {ph} ({cnt} طلب)")

    if report_type in ("general", "areas"):
        top = top_counts("area", 3)
        if top:
            lines.append("\n📍 *أفضل 3 مناطق:*")
            for i, (ar, cnt) in enumerate(top, 1):
                lines.append(f"  {i}. {ar} ({cnt} طلب)")

    if report_type in ("general", "items"):
        qty_by_item = {}
        for o in orders:
            item = o.get("item_name")
            if item in (None, ""):
                continue
            qty_by_item[item] = qty_by_item.get(item, 0) + _safe_int(o.get("quantity", 0))
        top = sorted(qty_by_item.items(), key=lambda kv: kv[1], reverse=True)[:3]
        if top:
            lines.append("\n💊 *أكثر 3 منتجات:*")
            medals = ["🥇", "🥈", "🥉"]
            for i, (it, qty) in enumerate(top):
                lines.append(f"  {medals[i]} {it} ({qty} وحدة)")

    pt_counts = top_counts("payment_type", None)
    if pt_counts:
        lines.append("\n💰 *نوع الدفع:*")
        for ptype, cnt in pt_counts:
            pct = round(cnt / len(orders) * 100)
            lines.append(f"  • {ptype}: {cnt} طلب ({pct}%)")

    top_dr = top_counts("doctor_name", 3)
    if top_dr:
        lines.append("\n🩺 *أكثر الأطباء:*")
        for i, (dr, cnt) in enumerate(top_dr, 1):
            lines.append(f"  {i}. {dr} ({cnt} طلب)")

    return "\n".join(lines)


def generate_template_excel(
    orders: list,
    catalog: list,
    rep_name: str,
    month: int = None,
    year: int = None,
    drug_targets: dict = None,
) -> io.BytesIO:
    """
    ينشئ Excel بنفس تنسيق قالب الوكيل:
    - الصفوف العلوية (2-7): ملخص الأدوية (TGT, ACCHIVMENT, price, units, SALES$)
    - الصف 8: ترويسة الأعمدة (صيدلية، منطقة، تاريخ، أدوية، API، %)
    - الصفوف 9+: زيارات الصيدليات (صف لكل زيارة، عمود لكل دواء)
    - آخر صف: معدل تحقيق التاركت
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "تقرير المبيعات"
    ws.sheet_view.rightToLeft = True

    today_d = date.today()
    m = month or today_d.month
    y = year or today_d.year
    tgt = drug_targets or {}

    drug_list = [d["drug_name"] for d in catalog] if catalog else []
    price_map = {d["drug_name"]: float(d.get("price") or 0) for d in catalog} if catalog else {}

    if not drug_list:
        seen_d = []
        seen_set = set()
        for o in orders:
            name = (o.get("item_name") or "").strip()
            if name and name not in seen_set:
                seen_set.add(name)
                seen_d.append(name)
        drug_list = seen_d

    n = len(drug_list)
    api_col = 4 + n
    pct_col = 5 + n

    fill_dark  = PatternFill("solid", fgColor="1F3864")
    fill_mid   = PatternFill("solid", fgColor="2E75B6")
    fill_yel   = PatternFill("solid", fgColor="FFF2CC")
    fill_grn   = PatternFill("solid", fgColor="E2EFDA")
    fill_org   = PatternFill("solid", fgColor="FCE4D6")
    fill_lgr   = PatternFill("solid", fgColor="F2F2F2")
    fill_blu   = PatternFill("solid", fgColor="BDD7EE")
    fill_row   = PatternFill("solid", fgColor="DEEAF1")

    fw  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    fb  = Font(name="Calibri", bold=True, size=10)
    fn  = Font(name="Calibri", size=10)
    cen = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rgt = Alignment(horizontal="right",  vertical="center", wrap_text=True)
    thn = Side(style="thin", color="AAAAAA")
    brd = Border(left=thn, right=thn, top=thn, bottom=thn)

    def S(cell, fill=None, font=None, align=None, fmt=None):
        try:
            from openpyxl.cell.cell import MergedCell as _MC
            if isinstance(cell, _MC):
                return
        except ImportError:
            pass
        if fill:  cell.fill = fill
        if font:  cell.font = font
        if align: cell.alignment = align
        if fmt:   cell.number_format = fmt
        cell.border = brd

    visit_keys = []
    seen_v = set()
    for o in sorted(orders, key=lambda x: x.get("order_date") or ""):
        k = (o.get("pharmacy") or "—", o.get("area") or "—", o.get("order_date") or "—")
        if k not in seen_v:
            seen_v.add(k)
            visit_keys.append(k)

    qty_map = {}
    for o in orders:
        k4 = (
            o.get("pharmacy") or "—",
            o.get("area") or "—",
            o.get("order_date") or "—",
            (o.get("item_name") or "").strip(),
        )
        qty_map[k4] = qty_map.get(k4, 0) + (o.get("quantity") or 0)
    qty_map_lower = {
        (ph, ar, dt, name.lower()): qty
        for (ph, ar, dt, name), qty in qty_map.items()
    }

    def _lookup_drug_units(catalog_name: str, ph: str, ar: str, dt: str) -> int:
        """بحث مرن عن الكمية — يجرب المطابقة الحرفية (case-insensitive) ثم جزئية ثم ضبابية."""
        cn = catalog_name.lower()
        # 1. مطابقة حرفية مباشرة (original case)
        exact = qty_map.get((ph, ar, dt, catalog_name), 0)
        if exact:
            return exact
        # 2. مطابقة حرفية case-insensitive
        exact_lower = qty_map_lower.get((ph, ar, dt, cn), 0)
        if exact_lower:
            return exact_lower
        # 3. مطابقة جزئية: اسم الكتالوج يبدأ باسم الدواء في الطلب أو العكس
        best = 0
        for (kph, kar, kdt, kname), qty in qty_map_lower.items():
            if kph != ph or kar != ar or kdt != dt:
                continue
            kn = kname.lower()
            if cn in kn or kn in cn or kn.startswith(cn[:5]):
                best += qty
        if best:
            return best
        # 4. مطابقة ضبابية بـ SequenceMatcher
        from difflib import SequenceMatcher
        best_score, best_qty = 0.0, 0
        for (kph, kar, kdt, kname), qty in qty_map_lower.items():
            if kph != ph or kar != ar or kdt != dt:
                continue
            score = SequenceMatcher(None, cn, kname.lower()).ratio()
            if score > best_score:
                best_score, best_qty = score, qty
        return best_qty if best_score >= 0.70 else 0

    drug_units = {d: sum(_lookup_drug_units(d, ph, ar, dt) for ph, ar, dt in visit_keys) for d in drug_list}
    drug_sales = {d: drug_units[d] * price_map.get(d, 0) for d in drug_list}
    total_tgt_u = sum(tgt.get(d, 0) for d in drug_list)
    total_units = sum(drug_units[d] for d in drug_list)
    total_sales = sum(drug_sales[d] for d in drug_list)
    total_ach   = round(total_units / total_tgt_u * 100, 1) if total_tgt_u > 0 else 0

    def ach_pct(drug):
        t = tgt.get(drug, 0)
        return round(drug_units[drug] / t * 100, 1) if t else None

    # ── Row 1: Title ──────────────────────────────────────────────────────────
    if n > 0:
        ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=api_col - 1)
    S(ws.cell(1, 4, f"ايتمات واجبة التحقيق — {rep_name} — {m}/{y}"), fill_dark, fw, cen)
    S(ws.cell(1, api_col, "TOTAL SALES"), fill_dark, fw, cen)
    ws.row_dimensions[1].height = 26

    # ── Row 2: ITEMS + drug names ──────────────────────────────────────────────
    S(ws.cell(2, 3, "ITEMS"), fill_dark, fw, cen)
    for i, d in enumerate(drug_list):
        S(ws.cell(2, 4 + i, d), fill_mid, fw, cen)
    ws.row_dimensions[2].height = 20

    # ── Row 3: TGT ────────────────────────────────────────────────────────────
    S(ws.cell(3, 3, "TGT"), fill_yel, fb, cen)
    for i, d in enumerate(drug_list):
        val = tgt.get(d) or ""
        S(ws.cell(3, 4 + i, val), fill_yel, fn, cen)
    S(ws.cell(3, api_col, total_tgt_u or ""), fill_yel, fb, cen)

    # ── Row 4: ACHIEVEMENT ────────────────────────────────────────────────────
    S(ws.cell(4, 3, "ACCHIVMENT"), fill_grn, fb, cen)
    for i, d in enumerate(drug_list):
        a = ach_pct(d)
        S(ws.cell(4, 4 + i, f"{a}%" if a is not None else ""), fill_grn, fn, cen)
    S(ws.cell(4, api_col, f"{total_ach}%" if total_ach else ""), fill_grn, fb, cen)

    # ── Row 5: price ──────────────────────────────────────────────────────────
    S(ws.cell(5, 3, "price"), fill_org, fb, cen)
    for i, d in enumerate(drug_list):
        p = price_map.get(d, 0)
        S(ws.cell(5, 4 + i, p or ""), fill_org, fn, cen, "#,##0")

    # ── Row 6: units ──────────────────────────────────────────────────────────
    S(ws.cell(6, 3, "units"), fill_lgr, fb, cen)
    for i, d in enumerate(drug_list):
        S(ws.cell(6, 4 + i, drug_units[d] or ""), fill_lgr, fn, cen)
    S(ws.cell(6, api_col, total_units or ""), fill_lgr, fb, cen)

    # ── Row 7: SALES $ ────────────────────────────────────────────────────────
    S(ws.cell(7, 3, "SALES $"), fill_blu, fb, cen)
    for i, d in enumerate(drug_list):
        s = drug_sales[d]
        S(ws.cell(7, 4 + i, round(s) if s else ""), fill_blu, fn, cen, "#,##0")
    S(ws.cell(7, api_col, round(total_sales) if total_sales else ""), fill_blu, fb, cen, "#,##0")

    # ── Row 8: Column headers ─────────────────────────────────────────────────
    for col, lbl in [(1, "الصيدلية"), (2, "المنطقة"), (3, "التاريخ")]:
        S(ws.cell(8, col, lbl), fill_dark, fw, cen)
    for i, d in enumerate(drug_list):
        S(ws.cell(8, 4 + i, d), fill_dark, fw, cen)
    S(ws.cell(8, api_col, "إجمالي"), fill_dark, fw, cen)
    S(ws.cell(8, pct_col, "%"), fill_dark, fw, cen)
    ws.row_dimensions[8].height = 20

    # ── Data rows (row 9+) ────────────────────────────────────────────────────
    for r_i, (ph, ar, dt) in enumerate(visit_keys):
        row = 9 + r_i
        rf = fill_row if r_i % 2 == 0 else None
        S(ws.cell(row, 1, ph), rf, fn, rgt)
        S(ws.cell(row, 2, ar), rf, fn, cen)
        # تنسيق التاريخ: نعرضه بصيغة "يوم/شهر" مطابقاً للقالب الأصلي
        try:
            dt_obj = datetime.strptime(dt, "%Y-%m-%d")
            display_dt = f"{dt_obj.day}/{dt_obj.month}"
        except Exception:
            display_dt = dt
        S(ws.cell(row, 3, display_dt), rf, fn, cen)
        row_u = 0
        row_s = 0.0
        for i, d in enumerate(drug_list):
            qty = _lookup_drug_units(d, ph, ar, dt)
            S(ws.cell(row, 4 + i, qty if qty else ""), rf, fn, cen)
            row_u += qty
            row_s += qty * price_map.get(d, 0)
        S(ws.cell(row, api_col, round(row_s) if row_s else ""), rf, fb, cen, "#,##0")
        if total_tgt_u > 0 and row_u > 0:
            S(ws.cell(row, pct_col, f"{round(row_u / total_tgt_u * 100, 1)}%"), rf, fn, cen)
        else:
            S(ws.cell(row, pct_col, ""), rf, fn, cen)

    # ── Final summary row ─────────────────────────────────────────────────────
    last = 9 + len(visit_keys)
    if n > 0:
        ws.merge_cells(start_row=last, start_column=1, end_row=last, end_column=api_col - 1)
    S(ws.cell(last, 1, "معدل تحقيق التاركت"), fill_dark, fw, cen)
    S(ws.cell(last, api_col, f"{total_ach}%" if total_ach else "—"), fill_dark, fw, cen)

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    for i in range(n):
        ws.column_dimensions[get_column_letter(4 + i)].width = 13
    if n > 0:
        ws.column_dimensions[get_column_letter(api_col)].width = 13
        ws.column_dimensions[get_column_letter(pct_col)].width = 8
    ws.freeze_panes = ws.cell(9, 4)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_admin_summary_excel(reps_data: list, month: int, year: int, achievement_map: dict = None) -> io.BytesIO:
    """ينشئ Excel يعرض ملخص مبيعات جميع المناديب للشهر المحدد مع نسبة التحقيق."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ملخص المناديب"
    ws.sheet_view.rightToLeft = True

    fill_hdr   = PatternFill("solid", fgColor="1F3864")
    fill_even  = PatternFill("solid", fgColor="DEEAF1")
    fill_tot   = PatternFill("solid", fgColor="FFF2CC")
    fill_grn   = PatternFill("solid", fgColor="E2EFDA")
    fill_red   = PatternFill("solid", fgColor="FCE4D6")
    fw   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    fb   = Font(name="Calibri", bold=True, size=10)
    fn   = Font(name="Calibri", size=10)
    cen  = Alignment(horizontal="center", vertical="center")
    rgt  = Alignment(horizontal="right",  vertical="center")
    thn  = Side(style="thin", color="AAAAAA")
    brd  = Border(left=thn, right=thn, top=thn, bottom=thn)

    def S(cell, fill=None, font=None, align=None, fmt=None):
        try:
            from openpyxl.cell.cell import MergedCell as _MC
            if isinstance(cell, _MC):
                return
        except ImportError:
            pass
        if fill:  cell.fill = fill
        if font:  cell.font = font
        if align: cell.alignment = align
        if fmt:   cell.number_format = fmt
        cell.border = brd

    months_ar = ["", "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
                 "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
    m_name = months_ar[month] if 1 <= month <= 12 else str(month)

    ws.merge_cells("A1:G1")
    S(ws.cell(1, 1, f"📊 ملخص مبيعات المناديب — {m_name} {year}"), fill_hdr, fw, cen)
    ws.row_dimensions[1].height = 28

    headers = ["#", "اسم المندوب", "طلبات", "كميات", "مبيعات (IQD)", "تحقيق التاركت", "الخطة"]
    for i, h in enumerate(headers, 1):
        S(ws.cell(2, i, h), fill_hdr, fw, cen)
    ws.row_dimensions[2].height = 20

    for r_i, rep in enumerate(reps_data):
        row = 3 + r_i
        ach = (achievement_map or {}).get(rep["id"], 0)
        rf = fill_grn if ach >= 80 else (fill_even if r_i % 2 == 0 else None)
        if 0 < ach < 50:
            rf = fill_red
        S(ws.cell(row, 1, r_i + 1), rf, fn, cen)
        S(ws.cell(row, 2, rep.get("name") or "مجهول"), rf, fb, rgt)
        S(ws.cell(row, 3, rep.get("order_count", 0)), rf, fn, cen)
        S(ws.cell(row, 4, rep.get("total_qty", 0)), rf, fn, cen)
        S(ws.cell(row, 5, rep.get("total_sales", 0) or 0), rf, fn, cen, "#,##0")
        S(ws.cell(row, 6, f"{ach}%"), rf, fb, cen)
        S(ws.cell(row, 7, rep.get("plan") or "—"), rf, fn, cen)

    total_row = 3 + len(reps_data)
    S(ws.cell(total_row, 1, ""), fill_tot, fb, cen)
    S(ws.cell(total_row, 2, "الإجمالي"), fill_tot, fb, rgt)
    S(ws.cell(total_row, 3, sum(r.get("order_count", 0) for r in reps_data)), fill_tot, fb, cen)
    S(ws.cell(total_row, 4, sum(r.get("total_qty", 0) for r in reps_data)), fill_tot, fb, cen)
    S(ws.cell(total_row, 5, sum(r.get("total_sales", 0) or 0 for r in reps_data)), fill_tot, fb, cen, "#,##0")
    S(ws.cell(total_row, 6, ""), fill_tot, fb, cen)
    S(ws.cell(total_row, 7, ""), fill_tot, fb, cen)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 14

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_compare_summary(orders1: list, orders2: list, name1: str, name2: str) -> str:
    def stats(orders):
        if not orders:
            return {"total": 0, "qty": 0, "top_ph": "—", "top_item": "—"}
        qty = sum(_safe_int(o.get("quantity", 0)) for o in orders)
        ph_counter = Counter(o.get("pharmacy") for o in orders if o.get("pharmacy") not in (None, ""))
        item_counter = Counter(o.get("item_name") for o in orders if o.get("item_name") not in (None, ""))
        return {
            "total": len(orders),
            "qty": qty,
            "top_ph": ph_counter.most_common(1)[0][0] if ph_counter else "—",
            "top_item": item_counter.most_common(1)[0][0] if item_counter else "—",
        }

    s1, s2 = stats(orders1), stats(orders2)
    diff_orders = s1["total"] - s2["total"]
    diff_qty = s1["qty"] - s2["qty"]
    arrow_o = "🔺" if diff_orders > 0 else ("🔻" if diff_orders < 0 else "➡️")
    arrow_q = "🔺" if diff_qty > 0 else ("🔻" if diff_qty < 0 else "➡️")

    return (
        f"📊 *مقارنة {name1} vs {name2}*\n\n"
        f"{'البيان':<20} {name1:<12} {name2}\n"
        f"{'─'*40}\n"
        f"عدد الطلبات:  {s1['total']}  ←→  {s2['total']}  {arrow_o} ({diff_orders:+d})\n"
        f"إجمالي الكمية: {s1['qty']}  ←→  {s2['qty']}  {arrow_q} ({diff_qty:+d})\n\n"
        f"🏆 أفضل صيدلية:\n  {name1}: {s1['top_ph']}\n  {name2}: {s2['top_ph']}\n\n"
        f"💊 أكثر منتج:\n  {name1}: {s1['top_item']}\n  {name2}: {s2['top_item']}"
    )



# ══════════════════════════════════════════════════════════════════════════════
# ── Survey Excel — 8-column canonical format (export + import) ───────────────
# ══════════════════════════════════════════════════════════════════════════════
#
# Column order (MUST stay stable — import depends on it):
#  1: المنطقة        (area)
#  2: اسم الطبيب    (name)
#  3: التخصص        (specialty)
#  4: أيام الدوام   (working_days)
#  5: الصيدلية المرتبطة (linked_pharmacy)
#  6: الكلاس        (class)
#  7: المنتج        (current_product)
#  8: ملاحظات       (notes)
#
# The "mapping layer" lives here: SURVEY_COLS.
# If the schema changes, update SURVEY_COLS and regenerate.
# ══════════════════════════════════════════════════════════════════════════════

SURVEY_COLS = [
    # (header_ar, db_field, editable_by_user, col_width)
    ("المنطقة",             "area",             True,  15),
    ("اسم الطبيب",          "name",             False, 26),
    ("التخصص",              "specialty",        False, 18),
    ("أيام الدوام",          "working_days",     True,  18),
    ("الصيدلية المرتبطة",   "linked_pharmacy",  True,  22),
    ("الكلاس",              "class",            False,  9),
    ("المنتج",              "current_product",  True,  22),
    ("ملاحظات",             "notes",            True,  26),
]
_SURVEY_NCOLS = len(SURVEY_COLS)
_SURVEY_HEADER_ROW = 3   # row 1 = title, row 2 = editable hint, row 3 = headers, 4+ = data


def generate_survey_excel(doctors: list, rep_name: str = "مندوب",
                          month: int = None, year: int = None) -> io.BytesIO:
    """
    Export doctor profiles as a survey Excel with the canonical 8-column format.
    Privacy: every rep's survey is scoped to their user_id — only the admin
    can request another rep's survey via the admin panel.
    The exported file can be re-imported via parse_survey_import_excel().
    """
    from openpyxl import Workbook
    today = date.today()
    m = month or today.month
    y = year or today.year

    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl not installed")

    MONTHS_AR = {
        1:"يناير",2:"فبراير",3:"مارس",4:"أبريل",5:"مايو",6:"يونيو",
        7:"يوليو",8:"أغسطس",9:"سبتمبر",10:"أكتوبر",11:"نوفمبر",12:"ديسمبر",
    }
    month_ar = MONTHS_AR.get(m, str(m))

    wb = Workbook()
    ws = wb.active
    ws.title = "استبيان_أطباء"
    ws.sheet_view.rightToLeft = True

    fill_title   = PatternFill("solid", fgColor="1F3864")
    fill_hint    = PatternFill("solid", fgColor="E7EFF8")
    fill_hdr     = PatternFill("solid", fgColor="2E75B6")
    fill_locked  = PatternFill("solid", fgColor="EDEDED")   # non-editable cells
    fill_a       = PatternFill("solid", fgColor="C6EFCE")
    fill_b       = PatternFill("solid", fgColor="FFEB9C")
    fill_c       = PatternFill("solid", fgColor="DEEAF1")
    fill_alt     = PatternFill("solid", fgColor="F5F9FD")

    fw   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    fh   = Font(name="Calibri", bold=True, color="1F3864", size=9)
    fb   = Font(name="Calibri", bold=True, size=10)
    fn   = Font(name="Calibri", size=10)
    fgrn = Font(name="Calibri", bold=True, color="375623", size=10)
    cen  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    rgt  = Alignment(horizontal="right",  vertical="center", wrap_text=True)
    thn  = Side(style="thin",  color="BBBBBB")
    thk  = Side(style="medium", color="2E75B6")
    brd  = Border(left=thn, right=thn, top=thn, bottom=thn)
    brd_thk = Border(left=thk, right=thk, top=thk, bottom=thk)

    def _fmt(cell, fill=None, font=None, align=None, bold_border=False):
        if fill:  cell.fill  = fill
        if font:  cell.font  = font
        if align: cell.alignment = align
        cell.border = brd_thk if bold_border else brd

    # ── Row 1: Title ────────────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=_SURVEY_NCOLS)
    tc = ws.cell(1, 1,
        f"📋 استبيان الأطباء الميداني  |  {rep_name}  |  {month_ar} {y}  |  {len(doctors)} طبيب")
    _fmt(tc, fill_title, fw, cen)
    ws.row_dimensions[1].height = 30

    # ── Row 2: Edit-hint bar ─────────────────────────────────────────────────
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=_SURVEY_NCOLS)
    hc = ws.cell(2, 1,
        "✏️  الأعمدة ذات الخلفية البيضاء قابلة للتعبئة ميدانياً  |"
        "  🔒 الأعمدة الرمادية تُحدَّث تلقائياً من البوت  |"
        "  بعد التعبئة ارفع الملف للبوت لاستيراده")
    _fmt(hc, fill_hint, fh, cen)
    ws.row_dimensions[2].height = 18

    # ── Row 3: Column headers ─────────────────────────────────────────────────
    for col, (hdr, field, editable, _) in enumerate(SURVEY_COLS, 1):
        c = ws.cell(3, col, hdr)
        _fmt(c, fill_hdr, fw, cen, bold_border=True)
    ws.row_dimensions[3].height = 22

    # ── Data rows ─────────────────────────────────────────────────────────────
    cls_fill = {"A": fill_a, "B": fill_b, "C": fill_c}
    for r_idx, doc in enumerate(doctors, 4):
        cls = (doc.get("class") or "C").upper()
        base_fill = cls_fill.get(cls, fill_alt if r_idx % 2 == 0 else None)

        for c_idx, (hdr, field, editable, _) in enumerate(SURVEY_COLS, 1):
            val = doc.get(field) or ""
            cell = ws.cell(r_idx, c_idx, val)
            cell_fill = fill_locked if not editable else (base_fill or fill_alt)
            cell_font  = fgrn if field == "class" else (fb if field == "name" else fn)
            cell_align = rgt if field in ("area", "name", "specialty", "linked_pharmacy", "notes") else cen
            _fmt(cell, cell_fill, cell_font, cell_align)

        ws.row_dimensions[r_idx].height = 18

    # Empty rows if no doctors
    if not doctors:
        for r in range(4, 14):
            for c_idx, (_, _, editable, _) in enumerate(SURVEY_COLS, 1):
                cell = ws.cell(r, c_idx, "")
                _fmt(cell, None if editable else fill_locked, fn, cen)

    # ── Column widths ─────────────────────────────────────────────────────────
    for col, (_, _, _, w) in enumerate(SURVEY_COLS, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = ws.cell(4, 1)

    # ── Sheet 2: Instructions + legend ───────────────────────────────────────
    ws2 = wb.create_sheet("تعليمات الاستيراد")
    ws2.sheet_view.rightToLeft = True
    instructions = [
        ("📋 تعليمات تعبئة الاستبيان واستيراده", True),
        ("", False),
        ("📌 خطوات الاستخدام:", True),
        ("  1. عبّئ الأعمدة ذات الخلفية البيضاء فقط خلال الزيارة الميدانية", False),
        ("  2. لا تغيّر أسماء الأعمدة أو ترتيبها", False),
        ("  3. لا تحذف أي صف من البيانات الموجودة", False),
        ('  4. ارفع الملف للبوت عبر زر "📥 استيراد استبيان"', False),
        ("  5. البوت سيحدّث بيانات الأطباء تلقائياً", False),
        ("", False),
        ("📊 شرح الأعمدة:", True),
    ]
    for hdr, field, editable, _ in SURVEY_COLS:
        lock_str = "🔒 يُحدَّث تلقائياً" if not editable else "✏️ قابل للتعبئة"
        instructions.append((f"  • {hdr}  ({lock_str})", False))
    instructions += [
        ("", False),
        ("🟢 A = 10+ طلبات مرتبطة  |  🟡 B = 5-9  |  🔵 C = أقل من 5", False),
        ("📧 للدعم: تواصل مع مدير البوت", False),
    ]
    for r, (txt, bold) in enumerate(instructions, 1):
        cell = ws2.cell(r, 1, txt)
        cell.font = Font(bold=bold, size=11, name="Calibri")
        cell.alignment = Alignment(horizontal="right")
    ws2.column_dimensions["A"].width = 72

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def parse_survey_import_excel(raw_bytes: bytes) -> list:
    """
    Parse a survey Excel file exported by generate_survey_excel().
    Returns a list of doctor dicts ready for upsert_doctor_profile().
    Validates column headers against SURVEY_COLS — raises ValueError if mismatched.
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)

    # Find the survey sheet (first sheet or "استبيان_أطباء")
    ws = None
    for sname in wb.sheetnames:
        if sname.startswith("استبيان") or sname == wb.sheetnames[0]:
            ws = wb[sname]
            break
    if ws is None:
        raise ValueError("لم يتم العثور على شيت الاستبيان")

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < _SURVEY_HEADER_ROW:
        raise ValueError("الملف لا يحتوي على بيانات كافية")

    # Row index 0-based: header row = _SURVEY_HEADER_ROW - 1 = 2
    header_row = rows[_SURVEY_HEADER_ROW - 1]

    # Build column index map from header names
    expected_headers = [col[0] for col in SURVEY_COLS]
    col_map = {}
    for c_idx, cell_val in enumerate(header_row):
        if cell_val:
            stripped = str(cell_val).strip()
            if stripped in expected_headers:
                field = SURVEY_COLS[expected_headers.index(stripped)][1]
                col_map[field] = c_idx

    if "name" not in col_map:
        raise ValueError(
            f"لم يتم التعرف على أعمدة الاستبيان.\n"
            f"تأكد أن الملف صادر من البوت ولم تتغير أسماء الأعمدة.\n"
            f"الأعمدة المطلوبة: {', '.join(expected_headers)}"
        )

    doctors = []
    for row in rows[_SURVEY_HEADER_ROW:]:   # data starts after header row
        if not row or not any(row):
            continue
        name = str(row[col_map["name"]]).strip() if row[col_map["name"]] else ""
        if not name or name.lower() == "none":
            continue
        doc = {}
        for field, c_idx in col_map.items():
            val = row[c_idx] if c_idx < len(row) else None
            doc[field] = str(val).strip() if val and str(val).strip() != "None" else None
        doctors.append(doc)

    wb.close()
    return doctors


def generate_full_backup_excel(data: dict) -> io.BytesIO:
    """يبني ملف Excel فيه شيت منفصل لكل قسم من بيانات البوت (نسخة احتياطية شاملة)."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)  # نحذف الشيت الافتراضي الفارغ، نضيف شيتاتنا بالأسفل
    any_sheet = False
    for section_name, rows in data.items():
        # اسم الشيت بتيليكرام/Excel محدود بـ31 حرف وما يقبل بعض الرموز
        sheet_name = section_name[:31]
        ws = wb.create_sheet(sheet_name)
        if rows:
            # union كل المفاتيح بترتيب أول ظهور — يطابق سلوك pandas.DataFrame(rows)
            # لما الصفوف تختلف بمفاتيحها (بعض الأعمدة ناقصة ببعض الصفوف)
            columns = []
            seen = set()
            for row in rows:
                for k in row:
                    if k not in seen:
                        seen.add(k)
                        columns.append(k)
            ws.append(columns)
            for row in rows:
                ws.append([row.get(c) for c in columns])
        else:
            ws.append(["ملاحظة"])
            ws.append(["لا يوجد بيانات"])
        any_sheet = True
    if not any_sheet:
        ws = wb.create_sheet("فارغ")
        ws.append(["ملاحظة"])
        ws.append(["لا يوجد بيانات"])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ─── التسديدات (تحصيل من الصيدليات) ─────────────────────────────────────────

def generate_payment_excel(payments: list, balances: list = None,
                            rep_totals: list = None, is_admin: bool = False) -> io.BytesIO:
    """
    ورقة 1 "التسديدات": كل عملية تسديد سطر (التاريخ، الصيدلية، المبلغ، طريقة الدفع، ملاحظة، [المندوب لو أدمن]).
    ورقة 2 "ملخص الصيدليات": من balances (نتيجة get_pharmacy_balance/get_rep_pharmacy_balances) —
        عمودا "إجمالي المبيعات" و"المتبقي" يظهران فقط لو فيه صيدلية عندها total_sales غير None.
    ورقة 3 "ملخص المندوبين" (is_admin فقط): من rep_totals.
    """
    from openpyxl import Workbook

    wb = Workbook()

    # ── ورقة 1: التسديدات ────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "التسديدات"
    ws1.sheet_view.rightToLeft = True

    cols = ["payment_date", "pharmacy_name", "amount", "payment_method", "notes"]
    headers = {
        "payment_date": "التاريخ", "pharmacy_name": "الصيدلية", "amount": "المبلغ",
        "payment_method": "طريقة الدفع", "notes": "ملاحظة", "rep_name": "المندوب",
    }
    if is_admin:
        cols = ["rep_name"] + cols
    ws1.append([headers[c] for c in cols])
    for p in (payments or []):
        ws1.append([p.get(c) for c in cols])

    total_payments = sum(_safe_int(p.get("amount")) for p in (payments or []))
    ws1.append([])
    ws1.append(["الإجمالي"] + [""] * (len(cols) - 2) + [total_payments])

    # ── ورقة 2: ملخص الصيدليات ───────────────────────────────────────────
    ws2 = wb.create_sheet("ملخص الصيدليات")
    balances = balances or []
    has_sales_data = any(b.get("total_sales") is not None for b in balances)

    if has_sales_data:
        ws2.append(["الصيدلية", "إجمالي التسديدات", "إجمالي المبيعات", "المتبقي"])
        for b in balances:
            ws2.append([
                b.get("pharmacy_name"),
                _safe_int(b.get("total_payments")),
                _safe_int(b.get("total_sales")) if b.get("total_sales") is not None else "—",
                _safe_int(b.get("remaining")) if b.get("remaining") is not None else "—",
            ])
    else:
        ws2.append(["الصيدلية", "إجمالي التسديدات"])
        ws2.append(["ملاحظة", "لا تتوفر أسعار بالكتالوج لحساب المتبقي"])
        for b in balances:
            ws2.append([b.get("pharmacy_name"), _safe_int(b.get("total_payments"))])

    # ── ورقة 3: ملخص المندوبين (أدمن فقط) ────────────────────────────────
    if is_admin and rep_totals:
        ws3 = wb.create_sheet("ملخص المندوبين")
        ws3.append(["المندوب", "إجمالي تسديداته"])
        for r in rep_totals:
            ws3.append([r.get("rep_name"), _safe_int(r.get("total"))])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
